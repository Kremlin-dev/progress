from flask import Flask, make_response, jsonify
from docx import Document
from docx2pdf import convert
import io
import os
import uuid
from datetime import datetime, timedelta

app = Flask(__name__)

# Ensure the existence of the temp directory
temp_dir = "./temp"
if not os.path.exists(temp_dir):
    os.makedirs(temp_dir)

def replace_placeholders(doc, replacements):
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            for key, value in replacements.items():
                placeholder = f'{{ {key} }}'
                if placeholder in run.text:
                    if isinstance(value, datetime):
                        # Format dates before replacing
                        formatted_value = value.strftime("%d %B, %Y")
                    else:
                        formatted_value = str(value)

                    run.text = run.text.replace(placeholder, formatted_value)

def generate_word_doc(replacements, template_path):
    doc = Document(template_path)
    replace_placeholders(doc, replacements)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer

@app.route("/generate_acceptance_letter/<uuid:user_id>", methods=["GET"])
def generate_acceptance_letter(user_id):
    try:
        # Check if the user_id is a valid UUID before proceeding
        try:
            user_id_uuid = uuid.UUID(str(user_id))
        except ValueError:
            return jsonify({"status": 400, "message": "Invalid UUID format"}), 400

        # Define the replacements dictionary
        replacements = {
    "name": "John Doe",
    "start_date": "1 October, 2023",
    "end_date": "1 November, 2023",
    "department": "Engineering",
    "location": "Accra",
    # Add more fields as needed
}


        template_path = "./letter/acceptance.docx"

        if not os.path.exists(template_path):
            return jsonify({"status": 404, "message": "Template not found"}), 404

        # Generate the Word document with replaced placeholders
        word_buffer = generate_word_doc(replacements, template_path)

        # Save the Word document to a temporary file
        temp_word_file = f"{temp_dir}/temp_doc_{user_id}.docx"
        with open(temp_word_file, "wb") as temp_file:
            temp_file.write(word_buffer.getvalue())

        # Convert the Word document to PDF
        temp_pdf_file = f"{temp_dir}/temp_doc_{user_id}.pdf"
        convert(temp_word_file, temp_pdf_file)

        # Send the PDF document as a response
        with open(temp_pdf_file, "rb") as pdf_file:
            pdf_buffer = io.BytesIO(pdf_file.read())

        response = make_response(pdf_buffer.read())
        response.headers["Content-Type"] = "application/pdf"
        response.headers["Content-Disposition"] = "inline; filename=acceptance_letter.pdf"

        # Clean up temporary files
        os.remove(temp_word_file)
        os.remove(temp_pdf_file)

        return response

    except Exception as e:
        print("Error:", str(e))
        return jsonify({"status": 500, "message": str(e)}), 500

if __name__ == "__main__":
    app.run(debug=True, host='0.0.0.0', port=5000, threaded=True)
